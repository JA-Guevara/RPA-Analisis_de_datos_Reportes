from __future__ import annotations

import sys
import unittest
from pathlib import Path

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import pandas as pd
from sqlalchemy import create_engine, text

from itacamba_reportes.core.audit.hashing import add_hashes_to_df, compute_row_hash
from itacamba_reportes.core.db.base import Base
from itacamba_reportes.core.db.bulk import bulk_insert_df
from itacamba_reportes.inputs.referencias.models import ChannelRefModel, RouteRefModel


class TestEltPipeline(unittest.TestCase):
    def test_audit_hashing(self):
        row1 = {"col_a": "123", "col_b": "ABC"}
        row2 = {"col_a": "123", "col_b": "ABC"}
        row3 = {"col_a": "123", "col_b": "XYZ"}

        h1 = compute_row_hash(row1, ("col_a", "col_b"))
        h2 = compute_row_hash(row2, ("col_a", "col_b"))
        h3 = compute_row_hash(row3, ("col_a", "col_b"))

        self.assertEqual(h1, h2)
        self.assertNotEqual(h1, h3)
        self.assertEqual(len(h1), 64)

    def test_add_hashes_to_df(self):
        df = pd.DataFrame({
            "src_1": ["A", "B"],
            "src_2": [1, 2],
            "der_1": ["X", "Y"],
        })
        df_hashed = add_hashes_to_df(df, ("src_1", "src_2"), ("der_1",))
        self.assertIn("source_hash", df_hashed.columns)
        self.assertIn("derived_hash", df_hashed.columns)
        self.assertIn("full_hash", df_hashed.columns)
        self.assertEqual(len(df_hashed), 2)

    def test_bulk_insert_with_engine(self):
        mem_engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=mem_engine)

        df = pd.DataFrame([
            {"descripcion_sap": "TEST_ALMACEN_1", "canal": "DIRECTO", "segmento": "RETAIL"},
            {"descripcion_sap": "TEST_ALMACEN_2", "canal": "DISTRIB", "segmento": "MAYORISTA"},
        ])
        rows = bulk_insert_df(ChannelRefModel, df, truncate_first=True, engine=mem_engine)
        self.assertEqual(rows, 2)

        with mem_engine.connect() as conn:
            count = conn.execute(
                text("SELECT count(*) FROM int_ref_almacen WHERE descripcion_sap LIKE 'TEST_ALMACEN_%'")
            ).scalar()
            self.assertEqual(count, 2)

    def test_route_ref_model_bulk_insert(self):
        mem_engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(bind=mem_engine)

        df = pd.DataFrame([
            {"codigo_ruta": "4B2910", "ruta": "RUTA_1", "departamento": "SANTA CRUZ", "localidad": "YANAIGUA"},
            {"codigo_ruta": "4B2911", "ruta": "RUTA_2", "departamento": "SANTA CRUZ", "localidad": "VILLA ESPERANZA"},
        ])
        rows = bulk_insert_df(RouteRefModel, df, truncate_first=True, engine=mem_engine)
        self.assertEqual(rows, 2)

        with mem_engine.connect() as conn:
            count = conn.execute(
                text("SELECT count(*) FROM int_ref_ruta WHERE codigo_ruta LIKE '4B291%'")
            ).scalar()
            self.assertEqual(count, 2)

    def test_sql_transformation_files_exist(self):
        src_root = Path(__file__).resolve().parents[1] / "src" / "itacamba_reportes"
        expected_sql = [
            src_root / "inputs" / "cartera" / "transform.sql",
            src_root / "inputs" / "query" / "transform.sql",
            src_root / "inputs" / "gd20" / "transform.sql",
            src_root / "outputs" / "despacho_bi" / "transform.sql",
            src_root / "outputs" / "despacho_mes" / "transform.sql",
        ]
        for p in expected_sql:
            self.assertTrue(p.exists(), f"Falta archivo SQL: {p}")
            content = p.read_text(encoding="utf-8").strip()
            self.assertGreater(len(content), 50, f"El archivo SQL {p.name} está vacío")

    def test_model_field_types(self):
        from itacamba_reportes.inputs.cartera.models import CustomerIntModel, CustomerStgModel
        from itacamba_reportes.inputs.query.models import QueryIntModel
        from itacamba_reportes.inputs.referencias.models import ChannelRefModel, RouteRefModel

        self.assertEqual(ChannelRefModel.segmento.type.length, 100)
        self.assertEqual(QueryIntModel.dia.type.python_type, int)
        self.assertEqual(QueryIntModel.despacho_t.type.precision, 18)

        stg_cols = {c.name for c in CustomerStgModel.__table__.columns}
        self.assertTrue({"nombre_cliente", "vendedor", "equipo", "region_nombre"}.issubset(stg_cols))

        unique_indexes = [idx for idx in CustomerIntModel.__table__.indexes if idx.unique]
        self.assertTrue(any("cliente" in [c.name for c in idx.columns] for idx in unique_indexes))

    def test_sap_rpa_module(self):
        from dataclasses import FrozenInstanceError

        from itacamba_reportes.sap import (
            SELECTORES_SAP,
            SapConnectionError,
            SapDownloadConfig,
            SapSessionManager,
            SelectoresGD20,
            SelectoresSAP,
            selectors,
        )
        from itacamba_reportes.sap.actions import (
            execute_download_gd20,
            execute_download_query,
            execute_download_report,
            execute_login,
        )

        self.assertTrue(hasattr(selectors, "OKCODE"))
        self.assertTrue(hasattr(selectors, "GD20_LEDGER"))
        self.assertTrue(hasattr(selectors, "GD20_MENU_EXPORT_EXCEL"))

        # Verificacion de Selectores Estructurados e Inmutables (Dependency Injection)
        self.assertEqual(SELECTORES_SAP.ventana.main, "wnd[0]")
        self.assertEqual(SELECTORES_SAP.barra.okcode, "wnd[0]/tbar[0]/okcd")
        self.assertEqual(SELECTORES_SAP.login.usuario, "wnd[0]/usr/txtRSYST-BNAME")
        self.assertEqual(SELECTORES_SAP.gd20.ledger, "wnd[0]/usr/ctxtRGUD0-LEDGER")
        self.assertEqual(SELECTORES_SAP.zbcs.okcode, "zbcs")

        # Inmutabilidad estricta (frozen=True)
        with self.assertRaises(FrozenInstanceError):
            SELECTORES_SAP.gd20.ledger = "MODIFICADO"

        # Inyeccion de dependencias con selectores customizados
        custom_selectors = SelectoresSAP(gd20=SelectoresGD20(ledger="CUSTOM_LEDGER_LOCATOR"))
        self.assertEqual(custom_selectors.gd20.ledger, "CUSTOM_LEDGER_LOCATOR")
        self.assertEqual(custom_selectors.barra.okcode, "wnd[0]/tbar[0]/okcd")

        mgr = SapSessionManager(connection_index=0, session_index=0)
        self.assertEqual(mgr.connection_index, 0)
        self.assertEqual(mgr.session_index, 0)

        cfg_dynamic = SapDownloadConfig(
            transaction="gd20",
            ledger="VC",
            variant="Variante_Cierre",
            period_low=10,
            period_high=10,
            export_filename="GD20_OCTUBRE",
        )
        self.assertEqual(cfg_dynamic.transaction, "gd20")
        self.assertEqual(cfg_dynamic.period_low, 10)
        self.assertEqual(cfg_dynamic.export_filename, "GD20_OCTUBRE")

        cfg_gd20 = SapDownloadConfig.for_gd20(transaction="gd20", period=9)
        self.assertEqual(cfg_gd20.period_low, 9)
        self.assertEqual(cfg_gd20.period_high, 9)
        self.assertEqual(cfg_gd20.ledger, "VC")

        cfg_query = SapDownloadConfig.for_query(transaction="zbcs", date_from="01.10.2026", date_to="15.10.2026")
        self.assertEqual(cfg_query.transaction, "zbcs")
        self.assertEqual(cfg_query.date_from, "01.10.2026")

    def test_defensive_reader_validation(self):
        import tempfile
        from itacamba_reportes.core.io.reader import (
            ExcelSourceSpec,
            FastExcelReader,
            MissingColumnsError,
            SheetNotFoundError,
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            test_file = tmp_path / "TEST_DATA.xlsx"
            df = pd.DataFrame({"col_a": [1, 2], "col_b": ["x", "y"]})
            with pd.ExcelWriter(test_file) as writer:
                df.to_excel(writer, sheet_name="HOJA_REAL", index=False)

            # Caso 1: Hoja inexistente lanza SheetNotFoundError de inmediato
            spec_bad_sheet = ExcelSourceSpec(
                name="test",
                file_pattern="TEST_DATA.xlsx",
                sheet="HOJA_INEXISTENTE",
            )
            reader_bad_sheet = FastExcelReader(spec_bad_sheet, tmp_path)
            with self.assertRaises(SheetNotFoundError):
                reader_bad_sheet.read_raw()

            # Caso 2: Columna requerida faltante lanza MissingColumnsError de inmediato
            spec_missing_col = ExcelSourceSpec(
                name="test",
                file_pattern="TEST_DATA.xlsx",
                sheet="HOJA_REAL",
                required_columns=("col_a", "col_c_obligatoria"),
            )
            reader_missing_col = FastExcelReader(spec_missing_col, tmp_path)
            with self.assertRaises(MissingColumnsError):
                reader_missing_col.read_raw()

            # Caso 3: Lectura exitosa con validacion aprobada
            spec_ok = ExcelSourceSpec(
                name="test",
                file_pattern="TEST_DATA.xlsx",
                sheet="HOJA_REAL",
                required_columns=("col_a", "col_b"),
            )
            reader_ok = FastExcelReader(spec_ok, tmp_path)
            df_read = reader_ok.read_raw()
            self.assertEqual(len(df_read), 2)
            self.assertIn("col_a", df_read.columns)

    def test_clean_exporter_excludes_technical_columns(self):
        import tempfile
        import openpyxl
        from itacamba_reportes.core.io.exporter import ExcelReportExporter

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            exporter = ExcelReportExporter(tmp_path)

            df_with_tech = pd.DataFrame({
                "cliente": [101, 102],
                "volumen": [15.5, 20.0],
                "id": [1, 2],
                "created_at": ["2026-10-22", "2026-10-22"],
                "version": [1, 2],
                "is_current": [True, True],
                "source_hash": ["hash1", "hash2"],
            })

            target = exporter.export(df_with_tech, "REPORTE_TEST")
            wb = openpyxl.load_workbook(target)

            # Verificacion 1: No existe la hoja de CONTROL por defecto
            self.assertNotIn("CONTROL", wb.sheetnames)
            self.assertIn("Data", wb.sheetnames)

            # Verificacion 2: No se exportan columnas tecnicas
            sheet = wb["Data"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("cliente", headers)
            self.assertIn("volumen", headers)
            self.assertNotIn("id", headers)
            self.assertNotIn("created_at", headers)
            self.assertNotIn("version", headers)
            self.assertNotIn("is_current", headers)
            self.assertNotIn("source_hash", headers)

    def test_exporter_multi_sheet_historial_ventas(self):
        import tempfile
        import openpyxl
        from itacamba_reportes.core.io.exporter import ExcelReportExporter

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            exporter = ExcelReportExporter(tmp_path)

            df_ventas = pd.DataFrame({
                "Nº documento": [1001, 1002],
                "Cantidad": [10.5, 20.0],
                "id": [1, 2],
                "is_current": [True, True],
            })
            df_despachos = pd.DataFrame({
                "Documento venta": [5001, 5002],
                "x": [15, 16],
                "created_at": ["2026-10-26", "2026-10-26"],
            })

            target = exporter.export_multi_sheet(
                sheets={
                    "VENTAS": df_ventas,
                    "DESPACHOS": df_despachos,
                },
                file_name="HISTORIAL VENTAS.xlsx",
                timestamped=False,
            )

            self.assertEqual(target.name, "HISTORIAL VENTAS.xlsx")
            wb = openpyxl.load_workbook(target)
            self.assertEqual(wb.sheetnames, ["VENTAS", "DESPACHOS"])

            ws_v = wb["VENTAS"]
            headers_v = [cell.value for cell in ws_v[1]]
            self.assertIn("Nº documento", headers_v)
            self.assertNotIn("id", headers_v)
            self.assertNotIn("is_current", headers_v)
            self.assertEqual(ws_v.cell(2, headers_v.index("Nº documento") + 1).value, 1001)
            self.assertEqual(ws_v.cell(2, headers_v.index("Cantidad") + 1).value, 10.5)

            ws_d = wb["DESPACHOS"]
            headers_d = [cell.value for cell in ws_d[1]]
            self.assertIn("Documento venta", headers_d)
            self.assertIn("x", headers_d)
            self.assertNotIn("created_at", headers_d)
            self.assertEqual(ws_d.cell(2, headers_d.index("Documento venta") + 1).value, 5001)
            self.assertEqual(ws_d.cell(2, headers_d.index("x") + 1).value, 15)
            wb.close()

    def test_exporter_single_sheet_values_not_blank(self):
        import tempfile
        import openpyxl
        from itacamba_reportes.core.io.exporter import ExcelReportExporter

        with tempfile.TemporaryDirectory() as tmp:
            exporter = ExcelReportExporter(Path(tmp))
            df = pd.DataFrame({
                "Nº documento": [1001, 1002],
                "Cantidad": [10.5, 20.0],
                "id": [1, 2],
            })
            target = exporter.export(df, "TEST_SINGLE")
            wb = openpyxl.load_workbook(target)
            ws = wb["Data"]
            headers = [cell.value for cell in ws[1]]
            self.assertEqual(ws.cell(2, headers.index("Nº documento") + 1).value, 1001)
            self.assertEqual(ws.cell(2, headers.index("Cantidad") + 1).value, 10.5)
            wb.close()

    def test_alembic_single_baseline_schema(self):
        from alembic.config import Config
        from alembic.script import ScriptDirectory
        from itacamba_reportes.core.settings import EnvConfig
        import subprocess

        alembic_versions_dir = EnvConfig.PROJECT_ROOT / "alembic" / "versions"
        py_files = [f for f in alembic_versions_dir.glob("*.py") if not f.name.startswith("__")]
        # Debe existir exactamente 1 archivo de migracion baseline inicial
        self.assertEqual(len(py_files), 1, f"Se esperaba exactamente 1 version de Alembic, encontradas: {py_files}")

        cfg = Config(str(EnvConfig.PROJECT_ROOT / "alembic.ini"))
        cfg.set_main_option("script_location", str(EnvConfig.PROJECT_ROOT / "alembic"))
        script_dir = ScriptDirectory.from_config(cfg)
        heads = script_dir.get_heads()
        self.assertEqual(len(heads), 1)
        self.assertEqual(heads[0], "0001_baseline")

        # Verificar que la generacion offline de SQL para upgrade head funciona sin errores
        res = subprocess.run(
            [sys.executable, "-m", "alembic", "upgrade", "head", "--sql"],
            cwd=str(EnvConfig.PROJECT_ROOT),
            capture_output=True,
            text=True,
        )
        self.assertEqual(res.returncode, 0, f"alembic upgrade head --sql fallo: {res.stderr}")
        self.assertIn("CREATE TABLE int_cartera", res.stdout)
        self.assertIn("CREATE TABLE int_ref_ruta", res.stdout)
        self.assertIn("CREATE TABLE int_ventas_bi", res.stdout)
        self.assertIn("CREATE TABLE pipeline_executions", res.stdout)
        self.assertIn("0001_baseline", res.stdout)

    def test_column_mappings_export(self):
        from itacamba_reportes.core.io.column_mappings import (
            VENTAS_BI_MAPPING,
            DESPACHO_MES_MAPPING,
            DESPACHO_BI_MAPPING,
            map_export_columns,
        )
        self.assertEqual(len(VENTAS_BI_MAPPING), 77)
        self.assertEqual(len(DESPACHO_MES_MAPPING), 20)
        self.assertEqual(len(DESPACHO_BI_MAPPING), 56)
        self.assertIn("Fecha", DESPACHO_MES_MAPPING.values())
        self.assertIn("Mes", DESPACHO_MES_MAPPING.values())
        self.assertIn("x", DESPACHO_BI_MAPPING.values())
        self.assertEqual(DESPACHO_BI_MAPPING["dia"], "x")
        self.assertEqual(list(DESPACHO_BI_MAPPING.values())[0], "Centro Proveedor")
        self.assertEqual(list(DESPACHO_BI_MAPPING.values())[-1], "x")

        df_dummy = pd.DataFrame({
            "gerencia": ["SCZ"],
            "equipo": ["EQ1"],
            "fecha": ["2026-10-01"],
            "mes": [10],
            "dia": [1],
        })
        mapped = map_export_columns(df_dummy, DESPACHO_MES_MAPPING)
        self.assertEqual(list(mapped.columns), ["GERENCIA", "EQUIPO", "Fecha", "Mes", "Día"])

        df_bi_dummy = pd.DataFrame({
            "centro_origen": ["4B05"],
            "dia": [20],
        })
        mapped_bi = map_export_columns(df_bi_dummy, DESPACHO_BI_MAPPING)
        self.assertEqual(list(mapped_bi.columns), ["Centro Proveedor", "x"])
        self.assertEqual(mapped_bi["x"].iloc[0], 20)

        # Validar que un DataFrame completo produce exactamente las 56 columnas oficiales en orden
        full_dummy = pd.DataFrame({k: [None] for k in DESPACHO_BI_MAPPING.keys()})
        mapped_full = map_export_columns(full_dummy, DESPACHO_BI_MAPPING)
        self.assertEqual(len(mapped_full.columns), 56)
        self.assertEqual(list(mapped_full.columns), list(DESPACHO_BI_MAPPING.values()))

    def test_migration_matches_all_models_exactly(self):
        import importlib.util
        from itacamba_reportes.core.db.base import Base
        import itacamba_reportes.core.db.schema

        spec = importlib.util.spec_from_file_location("baseline", "alembic/versions/0001_baseline_esquema_inicial_elt.py")
        baseline = importlib.util.module_from_spec(spec)

        created_tables = {}
        class MockOp:
            @staticmethod
            def create_table(name, *columns, **kwargs):
                created_tables[name] = [c.name for c in columns if getattr(c, "name", None)]
            @staticmethod
            def create_index(*args, **kwargs): pass
            @staticmethod
            def drop_table(*args, **kwargs): pass
            @staticmethod
            def drop_index(*args, **kwargs): pass
            @staticmethod
            def f(x): return x

        spec.loader.exec_module(baseline)
        baseline.op = MockOp()
        baseline.upgrade()

        self.assertEqual(len(created_tables), 18)
        self.assertEqual(len(Base.metadata.tables), 18)
        for tname, mtable in Base.metadata.tables.items():
            self.assertIn(tname, created_tables)
            m_cols = set(c.name for c in mtable.columns)
            mig_cols = set(created_tables[tname])
            self.assertEqual(m_cols, mig_cols, f"Discrepancia de columnas en tabla {tname}: {m_cols ^ mig_cols}")

    def test_ensure_all_tables_exist_creates_missing_table(self):
        from itacamba_reportes.core.db.base import Base
        from itacamba_reportes.core.db.schema import ensure_all_tables_exist
        import itacamba_reportes.core.db.connection as conn_mod
        from sqlalchemy import create_engine, inspect

        mem_eng = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(mem_eng, tables=[t for n, t in Base.metadata.tables.items() if n != "int_despacho_bi"])
        insp = inspect(mem_eng)
        self.assertFalse(insp.has_table("int_despacho_bi"))

        orig_engine = conn_mod._engine
        conn_mod.set_engine(mem_eng)
        try:
            ensure_all_tables_exist()
            insp2 = inspect(mem_eng)
            self.assertTrue(insp2.has_table("int_despacho_bi"))
            self.assertEqual(len(insp2.get_columns("int_despacho_bi")), 63)
        finally:
            conn_mod.set_engine(orig_engine)

    def test_zero_docstrings_in_src(self):
        import ast
        import glob
        src_path = Path(__file__).resolve().parents[1] / "src"
        docstring_findings = []
        for py_file in src_path.rglob("*.py"):
            with open(py_file, "r", encoding="utf-8") as f:
                tree = ast.parse(f.read(), filename=str(py_file))
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.ClassDef, ast.Module)):
                    ds = ast.get_docstring(node)
                    if ds:
                        docstring_findings.append(f"{py_file.name}:{node.name if hasattr(node, 'name') else 'module'}")
        self.assertEqual(docstring_findings, [])


if __name__ == "__main__":
    unittest.main()
