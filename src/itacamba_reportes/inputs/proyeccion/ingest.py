from __future__ import annotations

import logging
import re
from pathlib import Path
import openpyxl
import pandas as pd

from itacamba_reportes.core.db.bulk import bulk_insert_df
from itacamba_reportes.core.io.reader import MissingColumnsError, SheetNotFoundError
from itacamba_reportes.core.settings import EnvConfig
from itacamba_reportes.inputs.proyeccion.models import ProyeccionStgModel

logger = logging.getLogger(__name__)


def resolve_proyeccion_path() -> Path:
    candidates = [
        EnvConfig.DATA_PATH / "fijos" / "PROYECCION-SEMANAL.xlsx",
        EnvConfig.DATA_PATH / "PROYECCION-SEMANAL.xlsx",
        Path(r"C:\Users\joseg_iog5jdg\OneDrive - Tecorp\Documentos ITACAMBA\PROYECCION-SEMANAL.xlsx"),
        EnvConfig.PROJECT_ROOT / "storage" / "input" / "PROYECCION-SEMANAL.xlsx",
    ]
    for p in candidates:
        if p.exists():
            return p
    matches = list(EnvConfig.DATA_PATH.glob("*PROYECCION*SEMANAL*.xlsx")) if EnvConfig.DATA_PATH.exists() else []
    if matches:
        return matches[0]
    raise FileNotFoundError("No se encontró PROYECCION-SEMANAL.xlsx en las rutas configuradas")


def find_target_sheet(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=True)
    for sheet in wb.sheetnames:
        if re.search(r"^RF\s*AJUSTADO\s*SEMANAL", sheet, re.IGNORECASE):
            return sheet
    raise SheetNotFoundError(
        f"No se encontró una hoja que inicie con 'RF AJUSTADO SEMANAL' en '{path.name}'. Hojas disponibles: {wb.sheetnames}"
    )


def ingest_proyeccion() -> int:
    path = resolve_proyeccion_path()
    sheet_name = find_target_sheet(path)
    logger.info("Leyendo Proyección desde %s (hoja=%s)", path.name, sheet_name)

    try:
        df = pd.read_excel(path, sheet_name=sheet_name, engine="calamine")
    except Exception:
        df = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")

    renames = {
        "GERENCIA": "gerencia",
        "EQUIPO": "equipo",
        "DEPARTAMENTO": "departamento",
        "INCOTERMS": "incoterms",
        "CENTRO DISTRIBUCIÓN": "centro_distribucion",
        "CENTRO DISTRIBUCION": "centro_distribucion",
        "MATERIAL": "material",
        "LOCALIDAD IM": "localidad_im",
        "CODIGO CLIENTE": "codigo_cliente",
        "DIA": "dia",
        "RF": "rf",
        "RF AJUSTADO _ IM": "rf_ajustado_im",
        "RF AJUSTADO_IM": "rf_ajustado_im",
    }
    df = df.rename(columns=renames)

    # Validacion defensiva de columnas obligatorias
    required = ("dia", "rf_ajustado_im", "material")
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise MissingColumnsError(f"El archivo '{path.name}' (hoja '{sheet_name}') no contiene las columnas requeridas: {missing}")

    # Columna codigo_cliente: normalizacion numerica int64
    if "codigo_cliente" in df.columns:
        df["codigo_cliente"] = pd.to_numeric(df["codigo_cliente"], errors="coerce").fillna(0).astype("int64")

    if "dia" in df.columns:
        df["dia"] = pd.to_numeric(df["dia"], errors="coerce").fillna(0).astype(int)

    if "rf" in df.columns:
        df["rf"] = pd.to_numeric(df["rf"], errors="coerce").fillna(0)

    if "rf_ajustado_im" in df.columns:
        df["rf_ajustado_im"] = pd.to_numeric(df["rf_ajustado_im"], errors="coerce").fillna(0)

    model_cols = [c.name for c in ProyeccionStgModel.__table__.columns if c.name not in ("id", "created_at", "updated_at")]
    valid_cols = [c for c in model_cols if c in df.columns]

    rows = bulk_insert_df(ProyeccionStgModel, df[valid_cols], truncate_first=True)
    logger.info("Ingesta de Proyección a stg_proyeccion: %d filas", rows)
    return rows
